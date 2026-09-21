"""Deterministic DOCX report generator using python-docx."""
from __future__ import annotations


from datetime import date
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Inches, Pt, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from core.schemas import (
    StudentProfile,
    SubjectProfile,
    ExperimentDocument,
    GeneratedContent,
    ScreenshotArtifact,
)


class DocxGenerator:
    """Populate a DOCX template with experiment report content."""

    def __init__(self):
        self._insert_cursor = None

    def _add_paragraph(self, doc: Document):
        if self._insert_cursor is not None:
            return self._insert_cursor.insert_paragraph_before()
        return doc.add_paragraph()

    def generate(
        self,
        template_path: str,
        output_path: str,
        student: StudentProfile,
        subject: SubjectProfile,
        experiment: ExperimentDocument,
        generated: GeneratedContent,
        screenshots: dict[str, ScreenshotArtifact],
    ) -> str:
        """Generate a complete DOCX report and return the output path."""
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        # Copy template to preserve headers/footers/styles
        doc = Document(template_path)

        # Replace metadata in place to preserve styles, then clear old body
        self._replace_metadata(doc, student, subject, generated)
        self._clear_old_content(doc)

        # Build the report
        self._add_aim(doc, generated)

        # Add levels
        for level_name, label in [
            ("easy", "1. EASY LEVEL"),
            ("medium", "2. PROBLEM 2"),
            ("hard", "3. HARD LEVEL"),
        ]:
            level = getattr(generated, level_name, None)
            src_level = getattr(experiment, level_name, None)
            if level or src_level:
                self._add_level_section(
                    doc, level_name, label, level, src_level, screenshots
                )

        # Learning outcomes
        self._add_learning_outcomes(doc, generated)

        # Now delete the old content starting from the cursor
        self._delete_old_content(doc)

        try:
            doc.save(str(output))
            return str(output)
        except PermissionError:
            fallback = output.with_name(f"{output.stem}_v2{output.suffix}")
            doc.save(str(fallback))
            print(f"Warning: {output.name} was locked. Saved as {fallback.name} instead.")
            return str(fallback)

    def _add_custom_heading(self, doc: Document, text: str, level: int) -> None:
        """Add a heading using template styles. Only Heading 1 exists in the template."""
        p = self._add_paragraph(doc)
        if level == 1:
            try:
                p.style = "Heading 1"
            except KeyError:
                p.style = "Normal"
            p.paragraph_format.space_before = Emu(104775)
        else:
            # Sub-headings: use Body Text with bold since template has no Heading 2/3
            try:
                p.style = "Body Text"
            except KeyError:
                p.style = "Normal"
            p.paragraph_format.space_before = Emu(104775)
        run = p.add_run(text)
        run.bold = True
        return p

    def _replace_metadata(self, doc: Document, student: StudentProfile, subject: SubjectProfile, generated: GeneratedContent) -> None:
        """Replace placeholders in the template metadata header while preserving all original runs and styles."""
        today = date.today().strftime("%d/%m/%Y")
        title = generated.title or f"Experiment {generated.experiment_number}"
        
        replacements = {
            "Experiment 3": title,
            "{{STUDENT_NAME}}": student.student_name,
            "24BCS12733": student.uid,
            "BE-CSE": student.branch,
            "719-A": student.section_group,
            "{{SEMESTER}}": str(student.semester),
            "01/09/2026": today,
            "{{SUBJECT_NAME}}": subject.subject_name,
            "24CSH-301": subject.subject_code,
        }

        # The metadata is contained in the first 6 paragraphs.
        # Since MS Word splits words across runs, we merge runs for these paragraphs.
        # This is safe because all runs in these metadata paragraphs share the same styling (bold, font size).
        for i, p in enumerate(doc.paragraphs):
            if i > 5:
                break
            
            if not p.runs:
                continue

            # Reconstruct the full text of the paragraph
            full_text = "".join(r.text for r in p.runs)
            
            # Clean up zero-width characters that MS Word sometimes inserts
            full_text = full_text.replace('\xad', '')

            # Apply all replacements
            for old_text, new_text in replacements.items():
                full_text = full_text.replace(old_text, str(new_text))

            # Put the new text in the first run to preserve its styling
            p.runs[0].text = full_text

            # Clear the text of all subsequent runs to prevent duplication
            for r in p.runs[1:]:
                r.text = ""

    def _clear_old_content(self, doc: Document) -> None:
        """Locate the insertion cursor after the metadata (first 6 paragraphs)."""
        if len(doc.paragraphs) > 6:
            self._insert_cursor = doc.paragraphs[6]
        else:
            self._insert_cursor = None

    def _delete_old_content(self, doc: Document) -> None:
        """Remove all old paragraphs starting from the insertion cursor, preserving section breaks."""
        if self._insert_cursor is None:
            return
            
        body = doc.element.body
        to_remove = []
        found_cursor = False
        
        for child in list(body):
            if child == self._insert_cursor._element:
                found_cursor = True
                
            if found_cursor:
                if child.tag.endswith("}sectPr"):
                    continue
                    
                # PRESERVE paragraphs that contain section breaks to keep headers/footers/logos intact
                if child.xpath('.//w:sectPr'):
                    # Clear its runs so we don't leave old text, but keep the paragraph
                    for r in child.xpath('.//w:r'):
                        child.remove(r)
                    continue
                    
                to_remove.append(child)
                
        for child in to_remove:
            body.remove(child)

    def _add_aim(self, doc: Document, generated: GeneratedContent) -> None:
        aim = generated.global_aim or ""
        if aim:
            p = self._add_paragraph(doc)
            try:
                p.style = "Body Text"
            except KeyError:
                pass
            p.paragraph_format.space_before = Emu(204470)
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run("Aim: ")
            run.bold = True
            p.add_run(aim)

    def _add_level_section(
        self,
        doc: Document,
        level_name: str,
        label: str,
        generated_level,
        src_level,
        screenshots: dict[str, ScreenshotArtifact],
    ) -> None:
        """Add a complete level section (heading, problem, objectives, code, output)."""
        # Level heading
        self._add_custom_heading(doc, label, level=1)

        # Problem statement
        ps = ""
        if generated_level and generated_level.problem_statement:
            ps = generated_level.problem_statement
        elif src_level and src_level.problem_statement:
            ps = src_level.problem_statement
        if ps:
            p = self._add_paragraph(doc)
            try:
                p.style = "List Paragraph"
            except KeyError:
                pass
            p.paragraph_format.space_before = Emu(203835)
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run("Problem Statement: ")
            run.bold = True
            p.add_run(ps)

        # Objectives
        obj = ""
        if generated_level and generated_level.objectives:
            obj = generated_level.objectives
        if obj:
            p = self._add_paragraph(doc)
            try:
                p.style = "List Paragraph"
            except KeyError:
                pass
            p.paragraph_format.space_before = Emu(203835)
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run("Objective: ")
            run.bold = True
            p.add_run(obj)

        # Code
        code = ""
        if src_level and src_level.code:
            code = src_level.code
        elif generated_level and generated_level.code:
            code = generated_level.code

        if code:
            self._add_custom_heading(doc, "Code:", level=3)
            self._add_code_block(doc, code)

        # Output screenshot
        self._add_custom_heading(doc, "Output:", level=3)
        screenshot = screenshots.get(level_name)
        if screenshot and Path(screenshot.image_path).exists():
            p = self._add_paragraph(doc)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            try:
                run.add_picture(screenshot.image_path, width=Inches(5.5))
            except Exception:
                run.add_text("[Screenshot could not be inserted]")
        else:
            p = self._add_paragraph(doc)
            run = p.add_run("[Output screenshot not available]")
            run.font.italic = True

    def _add_code_block(self, doc: Document, code: str) -> None:
        """Add Java code as styled paragraphs."""
        import re
        code = re.sub(r'\n\s*\{', ' {', code)
        code = re.sub(r'\n\s*\[', ' [', code)
        code = re.sub(r'\n\s*\(', ' (', code)
        code = re.sub(r'\n\s*\n', '\n', code)

        lines = code.split("\n")
        for line in lines:
            p = self._add_paragraph(doc)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(11)
            run = p.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(10)

    def _add_learning_outcomes(
        self, doc: Document, generated: GeneratedContent
    ) -> None:
        """Add learning outcomes section."""
        self._add_custom_heading(doc, "5. Learning Outcome:", level=2)
        for outcome in generated.learning_outcomes:
            p = self._add_paragraph(doc)
            try:
                p.style = "Normal"
            except KeyError:
                pass
            p.add_run(f"\u2022 {outcome}")
