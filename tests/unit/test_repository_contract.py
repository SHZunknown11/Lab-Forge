from __future__ import annotations

import hashlib
import tomllib
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_required_directories_exist() -> None:
    required_directories = [
        ".agents/rules",
        "apps/api",
        "apps/web",
        "core/ingestion",
        "core/ai",
        "core/execution",
        "core/screenshots",
        "core/documents",
        "core/profiles",
        "core/validation",
        "core/pipeline",
        "templates/chandigarh-university/pblj",
        "examples",
        "artifacts",
        "tests/unit",
        "tests/integration",
        "tests/fixtures",
        "docs",
        "config",
    ]

    missing = [path for path in required_directories if not (ROOT / path).is_dir()]
    assert missing == []


def test_reference_assets_and_template_copy_exist() -> None:
    examples = ROOT / "examples"
    sample_docx = examples / "sample-report.docx"
    template_docx = ROOT / "templates/chandigarh-university/pblj/template.docx"

    assert (examples / "experiment-source.txt").exists()
    assert (examples / "source-notes.md").stat().st_size > 0
    assert sample_docx.stat().st_size > 0
    assert (examples / "sample-report.pdf").stat().st_size > 0
    assert template_docx.stat().st_size > 0

    # The sample-report.docx reference must remain unmodified (never edited).
    # The working template.docx was intentionally tokenized by scripts/tokenize_template.py
    # so its SHA now differs from the sample — that is expected and correct.
    # We verify tokenization succeeded by checking {{TOKEN}} placeholders are present.
    from docx import Document as _Document
    doc = _Document(str(template_docx))
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "{{" in full_text, (
        "Template does not appear to be tokenized. "
        "Run: python scripts/tokenize_template.py"
    )


def test_pyproject_is_valid_toml() -> None:
    data = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))

    assert data["project"]["name"] == "labforge"
    assert "fastapi" in "\n".join(data["project"]["dependencies"]).lower()
